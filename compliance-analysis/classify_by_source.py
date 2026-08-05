import openpyxl
from collections import defaultdict

ENTITY_MASTER = {"EST_NAME","EST_ADDRESS","EMPLOYER_NAME","PAN","TAN","GST_NO","LIN","EPF_CODE",
"EPF_ESI_CODES","PT_REG_NO","LWF_REG_NO","CIN","ESIC_EMPLOYER_CODE","REG_NO_UNDER_CODE","ENTITY_TYPE",
"NATURE_OF_BUSINESS","NIC_CODE","AUTHORISED_PERSON","EMPLOYER_CONTACT","CORRESPONDENCE_ADDR",
"BANK_ACCOUNT","GEO_COORDINATES","BRANCH_STRUCTURE","COMMENCEMENT_DATE","SITE_LOCATION",
"MANAGER_DETAILS","SECTOR_BLOCK","GROUP_COMPANIES","HEADER_BLOCK_REPEAT"}

PAYROLL_HR = {"WAGE_COMPONENTS","EMP_NAME","ATTENDANCE_DAYS","UAN","ESIC_IP_NO","CONTRIBUTION_ASSESSED",
"TAX_DEDUCTED","HEADCOUNT_MATRIX","EXIT_CAUSE_DATE","EMP_DOJ","EMP_DECLARATIONS","EMP_DOB","EMP_SEX",
"EMP_DESIGNATION","EMP_ADDRESS","EMP_CONTACT","EMP_FATHER_SPOUSE","EMP_MARITAL_STATUS","AADHAAR","PF_NO",
"NOMINEE_DETAILS","SEPARATIONS","SEPARATION_REASON","SEPARATION_NOTICE","SERVICE_YEARS","SERVICE_CHANGE",
"EMP_PAY_HISTORY","WAGE_PERIOD","WORK_PERIODS","BONUS_PAID","MIGRANT_COUNT","FTE_COUNT","GIG_WORKER_COUNT",
"EMP_DEPT","EMP_RELIGION","EMP_SUPERANNUATION","EPS_NPS","HOURS_SHIFTS","MATERNITY_STATS",
"MATERNITY_EVENT_DATES","MATERNITY_PAYMENTS","XVII_AGGREGATES","MANDAYS_LOST"}

LEDGER = {"PAYMENT_REFS","CONTRIBUTION_PAID","CONTRACT_COMMERCIALS","CONTRACTOR_LIST","ANNUAL_TURNOVER",
"CESS_AMOUNT","CESS_PAYMENTS","FIN_STATEMENTS","PRODUCTION_DATA","CONSTRUCTION_COST","RELIEF_AMOUNTS",
"MONEY_CLAIM"}

SYSTEM = {"RETURN_PERIOD","ESIGN"}

def bucket(code):
    if code in ENTITY_MASTER: return "Entity master"
    if code in PAYROLL_HR:    return "Payroll / attendance"
    if code in LEDGER:        return "General ledger"
    if code in SYSTEM:        return "System-generated"
    return "NOT derivable from books"

wb = openpyxl.load_workbook("Master_Labour_Forms_Element_Repetition.xlsx", data_only=True)

def as_int(v):
    try: return int(v)
    except: return 0

# annual transmission volume
fw = [r for r in list(wb["Frequency_Weighted"].iter_rows(values_only=True))[1:]
      if r[0] and len(str(r[0])) < 40]
vol = defaultdict(int); vol_n = defaultdict(int)
for r in fw:
    vol[bucket(r[0])] += as_int(r[2]); vol_n[bucket(r[0])] += 1
total_vol = sum(vol.values())

# raw ask counts
er = [r for r in list(wb["Element_Repetition"].iter_rows(values_only=True))[1:] if r[0]]
asks = defaultdict(int); els = defaultdict(int)
for r in er:
    asks[bucket(r[0])] += as_int(r[11]); els[bucket(r[0])] += 1
total_asks = sum(asks.values())

print("SHARE OF ANNUAL TRANSMISSIONS (574) BY SYSTEM OF RECORD")
for b in sorted(vol, key=lambda x: -vol[x]):
    print(f"  {b:<28} {vol[b]:>4}  ({100*vol[b]/total_vol:5.1f}%)   elements={vol_n[b]}")
derivable = total_vol - vol["NOT derivable from books"]
print(f"  --> derivable from books+payroll+master: {derivable}/{total_vol} = {100*derivable/total_vol:.1f}%")

print("\nSHARE OF RAW ASKS (376) BY SYSTEM OF RECORD")
for b in sorted(asks, key=lambda x: -asks[x]):
    print(f"  {b:<28} {asks[b]:>4}  ({100*asks[b]/total_asks:5.1f}%)   elements={els[b]}")
d2 = total_asks - asks["NOT derivable from books"]
print(f"  --> derivable: {d2}/{total_asks} = {100*d2/total_asks:.1f}%")

print("\nTHE NON-DERIVABLE RESIDUAL")
nd = [(r[0], as_int(r[11])) for r in er if bucket(r[0]) == "NOT derivable from books"]
nd.sort(key=lambda x: -x[1])
print(f"  {len(nd)} elements, {sum(k for _,k in nd)} asks")
print("  k>1:", ", ".join(f"{e}({k})" for e, k in nd if k > 1) or "none")
print(f"  singletons among them: {sum(1 for _,k in nd if k==1)}")
